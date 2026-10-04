# -*- coding: utf-8 -*-
"""A2 portrait editable conference poster for Symptom Tracker (LiA)."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Mm, Pt
from pathlib import Path

OUT = Path(__file__).with_name("Symptom_Tracker_LiA_Poster_A2.pptx")
QR = Path(__file__).with_name("github_qr.png")

# Teal palette
TEAL_800 = RGBColor(0x11, 0x5E, 0x59)
TEAL_700 = RGBColor(0x0F, 0x76, 0x6E)
TEAL_600 = RGBColor(0x0D, 0x94, 0x88)
TEAL_400 = RGBColor(0x2D, 0xD4, 0xBF)
TEAL_100 = RGBColor(0xCC, 0xFB, 0xF1)
TEAL_50 = RGBColor(0xF0, 0xFD, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x13, 0x4E, 0x4A)
SLATE = RGBColor(0x33, 0x41, 0x55)
PLACE = RGBColor(0x5E, 0xE0, 0xD0)
RULE = RGBColor(0x99, 0xF6, 0xE4)

W, H = Mm(420), Mm(594)
MARGIN = Mm(10)
GAP = Mm(6)
HEADER_H = Mm(54)
FOOTER_H = Mm(36)
COL_W = (W - MARGIN * 2 - GAP * 2) / 3
BODY_TOP = HEADER_H + Mm(8)
BODY_BOT = H - FOOTER_H - Mm(6)
BODY_H = BODY_BOT - BODY_TOP


def set_run_font(run, size, bold, color, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def stroke_shape(shape, color, pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def add_rect(slide, l, t, w, h, fill, line=None, line_pt=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill_shape(s, fill)
    if line is None:
        s.line.fill.background()
    else:
        stroke_shape(s, line, line_pt)
    s.shadow.inherit = False
    return s


def add_round(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_shape(s, fill)
    if line is None:
        s.line.fill.background()
    else:
        stroke_shape(s, line, 1.0)
    # tighter corner
    s.adjustments[0] = 0.06
    s.shadow.inherit = False
    return s


def add_tb(slide, l, t, w, h, paragraphs, valign=MSO_ANCHOR.TOP):
    """paragraphs: list of (text, size, bold, color, align, space_after)."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Mm(1.6)
    tf.margin_right = Mm(1.6)
    tf.margin_top = Mm(0.8)
    tf.margin_bottom = Mm(0.8)
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[valign],
        )
    except Exception:
        pass

    for i, item in enumerate(paragraphs):
        text, size, bold, color, align, after = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.clear()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = text
        set_run_font(run, size, bold, color)
    return box


def P(text, size=11, bold=False, color=INK, align=PP_ALIGN.LEFT, after=4):
    return (text, size, bold, color, align, after)


def add_mixed(slide, l, t, w, h, blocks, valign=MSO_ANCHOR.TOP):
    """blocks: list of list of (text, size, bold, color) for one paragraph, plus after/align."""
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Mm(2.4)
    tf.margin_right = Mm(2.4)
    tf.margin_top = Mm(2.0)
    tf.margin_bottom = Mm(2.0)
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr"}[valign],
        )
    except Exception:
        pass

    first = True
    for block in blocks:
        runs, after, align = block
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.clear()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(after)
        p.line_spacing = 1.12
        for text, size, bold, color in runs:
            run = p.add_run()
            run.text = text
            set_run_font(run, size, bold, color)
    return box


def R(text, size=11, bold=False, color=INK):
    return (text, size, bold, color)


def B(runs, after=5, align=PP_ALIGN.LEFT):
    return (runs, after, align)


def section_bar(slide, l, t, w, title):
    add_rect(slide, l, t, w, Mm(8.5), TEAL_700)
    add_tb(
        slide,
        l + Mm(2),
        t + Mm(0.6),
        w - Mm(4),
        Mm(7.4),
        [P(title, 13, True, WHITE, PP_ALIGN.LEFT, 0)],
        MSO_ANCHOR.MIDDLE,
    )
    return t + Mm(8.5)


def card(slide, l, t, w, h):
    return add_round(slide, l, t, w, h, TEAL_50, RULE)


def placeholder(slide, l, t, w, h, label):
    s = add_round(slide, l, t, w, h, TEAL_100, TEAL_600)
    add_tb(
        slide,
        l + Mm(3),
        t + (h - Mm(14)) / 2,
        w - Mm(6),
        Mm(14),
        [P(label, 11, True, TEAL_700, PP_ALIGN.CENTER, 0)],
        MSO_ANCHOR.MIDDLE,
    )
    return s


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_rect(slide, 0, 0, W, H, WHITE)
    add_rect(slide, 0, 0, W, HEADER_H, TEAL_700)
    add_rect(slide, 0, HEADER_H, W, Mm(2.2), TEAL_400)

    # Header text
    add_tb(
        slide,
        MARGIN,
        Mm(4),
        Mm(318),
        Mm(16),
        [P("SYMPTOM TRACKER", 32, True, WHITE, PP_ALIGN.LEFT, 0)],
    )
    add_tb(
        slide,
        MARGIN,
        Mm(20),
        Mm(318),
        Mm(10),
        [P(
            "A co-produced portable pain-event recorder for accessible symptom communication",
            13,
            False,
            TEAL_100,
            PP_ALIGN.LEFT,
            0,
        )],
    )
    add_tb(
        slide,
        MARGIN,
        Mm(31),
        Mm(318),
        Mm(9),
        [P(
            "Chuqing (Tracy) Tang  |  Department of Engineering, University of Cambridge  |  Murray Edwards College",
            11,
            True,
            WHITE,
            PP_ALIGN.LEFT,
            0,
        )],
    )
    add_tb(
        slide,
        MARGIN,
        Mm(41),
        Mm(318),
        Mm(10),
        [P(
            "Leadership in Action  |  July-August 2026  |  Supervisor: Vicky Gildersleve (Papworth Trust Cambridge Manager)",
            11,
            False,
            TEAL_100,
            PP_ALIGN.LEFT,
            0,
        )],
    )
    placeholder(slide, Mm(334), Mm(6), Mm(36), Mm(20), "Laidlaw logo")
    placeholder(slide, Mm(374), Mm(6), Mm(36), Mm(20), "Papworth logo")
    add_tb(
        slide,
        Mm(334),
        Mm(28),
        Mm(76),
        Mm(22),
        [
            P("LiA poster  |  SMART outcomes", 10, True, WHITE, PP_ALIGN.RIGHT, 1),
            P("UNESCO SDG 10  |  Reduced Inequalities", 10, False, TEAL_100, PP_ALIGN.RIGHT, 0),
        ],
    )

    c1 = MARGIN
    c2 = MARGIN + COL_W + GAP
    c3 = MARGIN + (COL_W + GAP) * 2

    # -------- Column 1 --------
    y = BODY_TOP
    y = section_bar(slide, c1, y, COL_W, "1   NEED AND USER RESEARCH")
    card(slide, c1, y, COL_W, Mm(78))
    add_mixed(
        slide,
        c1,
        y,
        COL_W,
        Mm(78),
        [
            B([R("Context (Papworth Trust co-production)", 12, True, TEAL_800)], 3),
            B([R(
                "Episodic pain is hard to remember and describe later. For many adults with learning disabilities, opening a dense health app and typing during an episode is unrealistic.",
                11, False, SLATE,
            )], 6),
            B([R("Discovery method (Weeks 1-2)", 12, True, TEAL_800)], 3),
            B([R("- Joined routine activities; observed recurring support needs", 11)], 2),
            B([R("- Spoke with service users, staff, and caregivers", 11)], 2),
            B([R("- Field checklist: independence, participation, confidence, safety, staff workload", 11)], 2),
            B([R("- Separated isolated incidents from recurring patterns", 11)], 5),
            B([R("Design brief: ", 11, True, TEAL_800), R(
                "record the moment with a physical squeeze; review it digitally later. Reduce reliance on memory, language load, and complex UI.",
                11, False, SLATE,
            )], 0),
        ],
    )
    y += Mm(80)
    placeholder(slide, c1, y, COL_W, Mm(38), "IMAGE: field checklist / needs research")
    y += Mm(41)

    y = section_bar(slide, c1, y, COL_W, "2   ENGINEERING APPROACH")
    h2 = BODY_BOT - y
    card(slide, c1, y, COL_W, h2)
    add_mixed(
        slide,
        c1,
        y,
        COL_W,
        h2,
        [
            B([R("User research set the architecture, not the reverse.", 11, True, TEAL_800)], 6),
            B([R("Sensing and event model (ESP32)", 12, True, TEAL_800)], 3),
            B([R("- FSR406 to ADC, with debounce and a press / release state machine", 11)], 2),
            B([R("- Stores a pressure-time curve, not only a peak value", 11)], 2),
            B([R("- Non-blocking haptic PWM so motor intensity follows force", 11)], 2),
            B([R("- Offline event queue on device storage when BLE is absent", 11)], 5),
            B([R("BLE sync (final solution)", 12, True, TEAL_800)], 3),
            B([R(
                "Queued episodes transfer on reconnect. The app claims each device event key before insert, so one press becomes one timeline row (no duplicate ends).",
                11, False, SLATE,
            )], 5),
            B([R("Calibration (person-specific, practical)", 12, True, TEAL_800)], 3),
            B([R(
                "A three-step weak / medium / strong map turns raw ADC into 0-100% relative force for that person, so readings stay meaningful across grip strength and mounting differences.",
                11, False, SLATE,
            )], 5),
            B([R("Hardware path", 12, True, TEAL_800)], 3),
            B([R(
                "Breadboard (LED stand-in for motor) to stripboard (driver, diode, power integrity) to SolidWorks enclosure shaped around thumb access and distinct SOS / REC controls.",
                11, False, SLATE,
            )], 0),
        ],
    )

    # -------- Column 2 --------
    y = BODY_TOP
    y = section_bar(slide, c2, y, COL_W, "3   SYSTEM OVERVIEW")
    placeholder(slide, c2, y, COL_W, Mm(48), "IMAGE: architecture  device / BLE / Flutter app")
    y += Mm(50)

    y = section_bar(slide, c2, y, COL_W, "4   HARDWARE AND IMPLEMENTATION")
    card(slide, c2, y, COL_W, Mm(62))
    add_mixed(
        slide,
        c2,
        y,
        COL_W,
        Mm(62),
        [
            B([R(
                "Stack: ESP32, FSR406, vibration motor (transistor + flyback), SOS / REC, status LED, optional INMP441 voice note, TP4056 battery path.",
                11, False, SLATE,
            )], 5),
            B([R("Integration lessons (engineering validity)", 12, True, TEAL_800)], 3),
            B([R("- Stripboard copper shorts can couple SOS and REC despite correct firmware", 11)], 2),
            B([R("- Motor load needs proper drive and bulk capacitance (battery brownout)", 11)], 2),
            B([R("- Silent WAV files: intermittent I2S clock / microphone slot config", 11)], 2),
            B([R("- Enclosure deformation can open joints: treat mechanics and electronics as one system", 11)], 0),
        ],
    )
    y += Mm(64)
    placeholder(slide, c2, y, COL_W, Mm(42), "IMAGE: stripboard / assembled device / CAD")
    y += Mm(45)

    y = section_bar(slide, c2, y, COL_W, "5   ACCESSIBLE MODE")
    text_h = Mm(72)
    card(slide, c2, y, COL_W, text_h)
    add_mixed(
        slide,
        c2,
        y,
        COL_W,
        text_h,
        [
            B([R(
                "Designed from centre feedback, not added as a theme toggle afterwards.",
                11, True, TEAL_800,
            )], 5),
            B([R("Service-user layer", 12, True, TEAL_800)], 3),
            B([R("- Larger targets and type; Home and Timeline only", 11)], 2),
            B([R("- 2 x 2 action grid; four picture tags (meal, exercise, sleep, medication)", 11)], 2),
            B([R("- SOS is a long-press, to reduce accidental alerts", 11)], 5),
            B([R("Caregiver layer", 12, True, TEAL_800)], 3),
            B([R(
                "Full settings remain available: calibration, vibration modes, curves, stats, labels, and Guardian / SOS options.",
                11, False, SLATE,
            )], 5),
            B([R("Principle: ", 11, True, TEAL_800), R(
                "people should not have to adapt themselves to an interface that was never designed around their needs.",
                11, False, SLATE,
            )], 0),
        ],
    )
    y += text_h + Mm(2)
    placeholder(slide, c2, y, COL_W, BODY_BOT - y, "IMAGE: Accessible Mode UI (Home grid / tags)")

    # -------- Column 3 --------
    y = BODY_TOP
    y = section_bar(slide, c3, y, COL_W, "6   EVALUATION (USABILITY)")
    card(slide, c3, y, COL_W, Mm(68))
    add_mixed(
        slide,
        c3,
        y,
        COL_W,
        Mm(68),
        [
            B([R(
                "Scope: usability and engineering validation. No clinical-efficacy claims.",
                11, True, TEAL_800,
            )], 4),
            B([R("Protocol", 12, True, TEAL_800)], 3),
            B([R("- Day 1 formative: observe first; explain only when needed; log critical / major issues", 11)], 2),
            B([R("- Prioritise fixes with a rationale; avoid adding features for every comment", 11)], 2),
            B([R("- Day 2: repeat the same core tasks and check clarity / reliability", 11)], 4),
            B([R("Questions: ", 11, True, TEAL_800), R(
                "minimal instruction? light vs strong presses? short vs long events? BLE workflow? can staff read the history?",
                11, False, SLATE,
            )], 3),
            B([R("Observed split: ", 11, True, TEAL_800), R(
                "users needed large targets, few choices, and physical feedback; staff needed time, duration, intensity, patterns, and one trustworthy row per press.",
                11, False, SLATE,
            )], 0),
        ],
    )
    y += Mm(70)
    placeholder(slide, c3, y, COL_W, Mm(36), "IMAGE: app UI / Accessible Mode / test checklist")
    y += Mm(39)

    y = section_bar(slide, c3, y, COL_W, "7   SMART GOALS TO OUTCOMES")
    card(slide, c3, y, COL_W, Mm(78))
    add_mixed(
        slide,
        c3,
        y,
        COL_W,
        Mm(78),
        [
            B([R("G1  Needs discovery (Weeks 1-2)", 12, True, TEAL_800)], 2),
            B([R(
                "S/M: field checklist and a written problem brief (physical record, digital review). Outcome: a focused TEAL brief on discomfort-communication barriers.",
                11, False, SLATE,
            )], 4),
            B([R("G2  Working prototype (Weeks 2-5)", 12, True, TEAL_800)], 2),
            B([R(
                "S/M: handheld ESP32 device and Flutter app with curve capture, offline queue, BLE sync, calibration, and Accessible Mode. Outcome: an end-to-end demo loop before the final week.",
                11, False, SLATE,
            )], 4),
            B([R("G3  Validation and sharing (Weeks 5-6)", 12, True, TEAL_800)], 2),
            B([R(
                "S/M: Day 1 / fix / Day 2 notes; LiA poster on the Network with SMART and SDG reporting. Outcome: documented feedback and an open repository.",
                11, False, SLATE,
            )], 4),
            B([R(
                "A / R / T: scoped to six weeks with Papworth priorities; milestones tracked from midpoint check-in to the final demo (August 2026).",
                11, True, TEAL_800,
            )], 0),
        ],
    )
    y += Mm(80)

    y = section_bar(slide, c3, y, COL_W, "8   UNESCO SDG 10  REDUCED INEQUALITIES")
    h8 = BODY_BOT - y
    card(slide, c3, y, COL_W, h8)
    placeholder(slide, c3 + Mm(3), y + Mm(4), Mm(22), Mm(22), "SDG 10")
    add_mixed(
        slide,
        c3 + Mm(26),
        y + Mm(2),
        COL_W - Mm(28),
        h8 - Mm(4),
        [
            B([R(
                "Inequality is not only whether a service exists, but whether people can access, understand, and benefit from it.",
                11, True, TEAL_800,
            )], 5),
            B([R(
                "Symptom Tracker targets barriers of memory, language load, and digital confidence in symptom communication, so support and review are more equitable between service users and caregivers.",
                11, False, SLATE,
            )], 0),
        ],
    )

    # Footer
    fy = H - FOOTER_H
    add_rect(slide, 0, fy, W, FOOTER_H, TEAL_800)
    add_tb(
        slide,
        MARGIN,
        fy + Mm(3),
        Mm(268),
        Mm(8),
        [P(
            "Co-produced with Papworth Trust  |  Murray Edwards College, University of Cambridge  |  Laidlaw Scholars",
            11,
            True,
            WHITE,
            PP_ALIGN.LEFT,
            0,
        )],
    )
    add_tb(
        slide,
        MARGIN,
        fy + Mm(12),
        Mm(268),
        Mm(20),
        [P(
            "Acknowledgements: service users, staff, and caregivers at Papworth Trust. Code: github.com/Tracytang-engi/Symptom_Tracker",
            10,
            False,
            TEAL_100,
            PP_ALIGN.LEFT,
            0,
        )],
    )
    placeholder(slide, Mm(282), fy + Mm(5), Mm(28), Mm(26), "Papworth")
    placeholder(slide, Mm(313), fy + Mm(5), Mm(28), Mm(26), "Laidlaw")
    placeholder(slide, Mm(344), fy + Mm(5), Mm(28), Mm(26), "Cambridge")

    if QR.exists():
        slide.shapes.add_picture(str(QR), Mm(380), fy + Mm(3), Mm(28), Mm(28))
        add_tb(
            slide,
            Mm(378),
            fy + Mm(30.5),
            Mm(32),
            Mm(5),
            [P("GitHub QR", 8, True, TEAL_400, PP_ALIGN.CENTER, 0)],
        )
    else:
        placeholder(slide, Mm(380), fy + Mm(4), Mm(28), Mm(28), "GitHub QR")

    prs.save(OUT)
    print("SAVED", OUT)


if __name__ == "__main__":
    build()
